from pathlib import Path

import markdown
from bs4 import BeautifulSoup
from docx import Document


INPUT_PATH = Path(__file__).parent / "SRS_Current_Implementation.md"
OUTPUT_PATH = Path(__file__).parent / "SRS_Current_Implementation.docx"


def text_of(node):
    return node.get_text(" ", strip=True)


def add_markdown_table(doc, table_tag):
    rows = table_tag.find_all("tr")
    if not rows:
        return

    max_cols = max(len(r.find_all(["th", "td"])) for r in rows)
    word_table = doc.add_table(rows=len(rows), cols=max_cols)
    word_table.style = "Table Grid"

    for r_idx, row in enumerate(rows):
        cells = row.find_all(["th", "td"])
        for c_idx, cell in enumerate(cells):
            word_table.cell(r_idx, c_idx).text = text_of(cell)


md_text = INPUT_PATH.read_text(encoding="utf-8")
html = markdown.markdown(md_text, extensions=["tables", "fenced_code", "sane_lists"])
soup = BeautifulSoup(html, "html.parser")


doc = Document()

for node in soup.children:
    if not getattr(node, "name", None):
        continue

    tag = node.name.lower()

    if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        level = int(tag[1])
        doc.add_heading(text_of(node), level=level)
    elif tag == "p":
        txt = text_of(node)
        if txt:
            doc.add_paragraph(txt)
    elif tag == "ul":
        for li in node.find_all("li", recursive=False):
            doc.add_paragraph(text_of(li), style="List Bullet")
    elif tag == "ol":
        for li in node.find_all("li", recursive=False):
            doc.add_paragraph(text_of(li), style="List Number")
    elif tag == "table":
        add_markdown_table(doc, node)
        doc.add_paragraph("")
    elif tag == "blockquote":
        txt = text_of(node)
        if txt:
            doc.add_paragraph(txt)
    elif tag == "pre":
        txt = text_of(node)
        if txt:
            doc.add_paragraph(txt)
    elif tag == "hr":
        doc.add_paragraph("----------------------------------------")


doc.save(OUTPUT_PATH)
print(f"Created: {OUTPUT_PATH}")
